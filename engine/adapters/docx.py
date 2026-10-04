from __future__ import annotations

from pathlib import Path
from docx import Document

from adapters.base import DocumentAdapter
from core.errors import EngineError


class DocxAdapter(DocumentAdapter):
    formats = ("docx",)
    media_types = ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",)

    def extract_text(self, path: str) -> str:
        try:
            doc = Document(path)
        except Exception as exc:
            raise EngineError("invalid_docx", "Unable to open DOCX document", 422, {"reason": str(exc)})
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(p for p in parts if p.strip())

    def validate(self, path: str) -> dict:
        text = self.extract_text(path)
        return {"valid": True, "format": "docx", "paragraphs": len(Document(path).paragraphs), "characters": len(text)}

    def write_from_text(self, text: str, path: str, options: dict) -> None:
        doc = Document()
        title = options.get("title")
        if title:
            doc.add_heading(str(title), level=1)
        for block in text.split("\n"):
            doc.add_paragraph(block)
        doc.save(path)
